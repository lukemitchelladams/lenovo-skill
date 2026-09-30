"""
Build a queryable knowledge base from the DCSC exports listed in dcsc_files.txt
(written by scan_dcsc.py).

Ground truth, not inference:
  - a config that contains BU1E is a DCSC-validated Top Choice Express build,
    so every feature code in it is TCE-compatible on that MTM
  - a config with no Critical message is a DCSC-validated buildable config
  - Message History rows are the real gating rules, verbatim

Incremental: existing configs are never deleted (their source files may be gone,
and _zipcache paths are renumbered every scan). A file is added only when its
content signature (export name + part lines) is not already in the db. The db is
never removed either, so the fact/doc_fts tables from build_kb_index.py survive
and a reader holding the file open (e.g. an MCP Sql tool) can't block the run.
"""
import os, re, sys, sqlite3, datetime, warnings, hashlib
import openpyxl
from paths import DB, KB_DIR

warnings.filterwarnings("ignore")
try:
    sys.stdout.reconfigure(errors="replace")
except Exception:
    pass

LIST = os.path.join(KB_DIR, "dcsc_files.txt")
if not os.path.exists(LIST):
    sys.exit("dcsc_files.txt not found - run first:  python scan_dcsc.py <folder holding your DCSC exports>")
FILES = [l.strip() for l in open(LIST, encoding="utf-8") if l.strip()]

os.makedirs(os.path.dirname(DB), exist_ok=True)
cx = sqlite3.connect(DB, timeout=60)
cx.executescript("""
CREATE TABLE IF NOT EXISTS config(
  id INTEGER PRIMARY KEY, file TEXT, name TEXT, mtime TEXT,
  mtm TEXT, model TEXT, nodes REAL, tce INTEGER, criticals INTEGER, total REAL);
CREATE TABLE IF NOT EXISTS part(
  cfg INTEGER, fc TEXT, descr TEXT, qty REAL, unit REAL, cat TEXT);
CREATE TABLE IF NOT EXISTS msg(
  cfg INTEGER, severity TEXT, text TEXT);
CREATE INDEX IF NOT EXISTS ix_part_fc ON part(fc);
CREATE INDEX IF NOT EXISTS ix_part_cfg ON part(cfg);
""")
if "sig" not in [r[1] for r in cx.execute("PRAGMA table_info(config)")]:
    cx.execute("ALTER TABLE config ADD COLUMN sig TEXT")
cx.execute("CREATE INDEX IF NOT EXISTS ix_config_sig ON config(sig)")

MTM_RE = re.compile(r"^7[A-Z0-9]{3}CTO[0-9A-Z]WW$")
ZIP_PREFIX = re.compile(r"^\d{4}_")

def export_name(path):
    # _zipcache copies carry a per-scan NNNN_ prefix; strip it so the name is stable
    base = os.path.basename(path)
    if "_zipcache" in path:
        base = ZIP_PREFIX.sub("", base)
    return base.lower()

def signature(path, parts):
    h = hashlib.sha1(export_name(path).encode("utf-8"))
    for fc, qty, unit in parts:
        h.update(f"\n{fc}\t{float(qty)!r}\t{float(unit)!r}".encode("utf-8"))
    return h.hexdigest()

# backfill signatures on rows built before incremental mode
for cid, f in cx.execute("SELECT id, file FROM config WHERE sig IS NULL").fetchall():
    rows = cx.execute("SELECT fc, qty, unit FROM part WHERE cfg=? ORDER BY rowid", (cid,)).fetchall()
    cx.execute("UPDATE config SET sig=? WHERE id=?", (signature(f, rows), cid))
known = {s for (s,) in cx.execute("SELECT sig FROM config")}
before = cx.execute("SELECT COUNT(*) FROM config").fetchone()[0]

def sheet_rows(ws):
    for row in ws.iter_rows(values_only=True):
        yield ["" if v is None else str(v).strip() for v in row]

ok = err = dup = 0
bad = []
for i, path in enumerate(FILES, 1):
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception as e:
        err += 1
        bad.append(f"{path}: {type(e).__name__}: {e}")
        continue
    names = wb.sheetnames
    qs = "Quote" if "Quote" in names else names[0]
    ws = wb[qs]

    parts, mtms, model, nodes, total, tce = [], [], "", None, 0.0, 0
    for c in sheet_rows(ws):
        if len(c) < 6:
            continue
        fc, descr = c[0], c[2]
        try:
            qty = float(c[4]); unit = float(c[5])
        except Exception:
            continue
        if not descr:
            continue
        total += qty * unit
        parts.append((fc, descr, qty, unit))
        if fc == "BU1E":
            tce = 1
        if MTM_RE.match(fc):
            mtms.append(fc)
            if ("Server" in descr or "Base Warranty" in descr) and not model:
                m = re.search(r"(ThinkSystem [A-Za-z0-9]+ ?V?\d?|ThinkAgile [A-Za-z0-9]+ ?V?\d?)", descr)
                model = m.group(1).strip() if m else descr[:40]
                nodes = qty
    mtm = max(set(mtms), key=mtms.count) if mtms else ""
    sig = signature(path, [(a, c_, d) for a, b, c_, d in parts])
    if sig in known:
        wb.close()
        dup += 1
        continue
    known.add(sig)

    msgs = []
    crit = 0
    if "Message History" in names:
        for c in sheet_rows(wb["Message History"]):
            if not c or not c[0] or c[0] == "Severity":
                continue
            sev = c[0]
            txt = c[1] if len(c) > 1 else ""
            if not txt:
                continue
            msgs.append((sev, txt[:600]))
            if sev.lower() == "critical":
                crit += 1
    wb.close()

    ts = datetime.datetime.fromtimestamp(os.path.getmtime(path)).strftime("%Y-%m-%d")
    cur = cx.execute(
        "INSERT INTO config(file,name,mtime,mtm,model,nodes,tce,criticals,total,sig) VALUES(?,?,?,?,?,?,?,?,?,?)",
        (path, os.path.basename(path), ts, mtm, model, nodes, tce, crit, total, sig))
    cid = cur.lastrowid
    cx.executemany("INSERT INTO part(cfg,fc,descr,qty,unit) VALUES(?,?,?,?,?)",
                   [(cid, a, b, c_, d) for a, b, c_, d in parts])
    cx.executemany("INSERT INTO msg(cfg,severity,text) VALUES(?,?,?)",
                   [(cid, s, t) for s, t in msgs])
    ok += 1
    print(f"  + {os.path.basename(path)}  tce={tce} crit={crit}")
    if i % 25 == 0:
        print(f"  ...{i}/{len(FILES)}")

cx.commit()

print(f"\nadded {ok} new configs ({dup} already in db, {err} unreadable)")
for b in bad:
    print(f"  unreadable: {b}")
print(f"configs      : {before} -> {cx.execute('SELECT COUNT(*) FROM config').fetchone()[0]}")
q = lambda s: cx.execute(s).fetchall()
print(f"parts rows   : {q('SELECT COUNT(*) FROM part')[0][0]:,}")
print(f"messages     : {q('SELECT COUNT(*) FROM msg')[0][0]:,}")
print(f"distinct MTMs: {q('SELECT COUNT(DISTINCT mtm) FROM config WHERE length(mtm)>0')[0][0]}")
print(f"TCE configs  : {q('SELECT COUNT(*) FROM config WHERE tce=1')[0][0]}")
print(f"with criticals: {q('SELECT COUNT(*) FROM config WHERE criticals>0')[0][0]}")

print("\n=== configs per MTM ===")
for mtm, model, n, t in q("""SELECT mtm, MIN(model), COUNT(*), SUM(tce)
                             FROM config WHERE mtm<>'' GROUP BY mtm ORDER BY COUNT(*) DESC"""):
    print(f"  {mtm}  {str(model)[:26]:<28} configs={n:<4} TCE-validated={t}")
print(f"\ndb: {DB}")
print(f"added {ok} new configs ({dup} already in db, {err} unreadable); total {cx.execute('SELECT COUNT(*) FROM config').fetchone()[0]}")
