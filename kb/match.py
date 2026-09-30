"""Competitor BOM -> Lenovo worksheet. Deterministic groundwork for a person or an AI to finish:
it never invents a feature code, it only proposes codes that exist in the DCSC rules snapshot for
the chosen MTM, ranked by how well they match each competitor line, with the TCE flag shown.

usage:
  python match.py <competitor-bom> [--platform "SR650 V4" | --mtm 7DGDCTO1WW] [--nodes N]
                                   [--tce] [--workload nutanix] [--json]

Competitor BOM: a text file (one item per line, e.g. '2 x Intel Xeon Gold 6526Y 16C 2.8GHz' or
'2<TAB>P52535-B21<TAB>Intel Xeon ...'), a .csv, or an .xlsx (the first sheet; a quantity column
and a description column are auto-detected). Quantities are per server unless the platform line's
quantity says otherwise (HPE and Dell exports carry order totals: see the note printed).

What it does:
  1. finds the competitor platform line and proposes Lenovo platforms (official Lenovo competitor
     map when kb/compete.py can reach it, plus form-factor / socket / CPU-vendor rules)
  2. parses every line (CPU, memory, drives, controller, boot, NIC, optics, PSU, GPU, FC, mgmt,
     services, software) and extracts the spec that matters
  3. proposes up to 3 Lenovo feature codes per line from the COMPLETE option list of the MTM
  4. runs the rule validators (rules.py) on the first-pick BOM and lists forced companions
Everything is a proposal. Build it in DCSC; only an export with BU1E proves Top Choice Express.
"""
import csv, json, os, re, sys

HERE = os.path.dirname(os.path.abspath(__file__))
if HERE not in sys.path:
    sys.path.insert(0, HERE)
import dcsc_rules

I = re.I

# ------------------------------------------------------------------ loading
def _num(x):
    try:
        v = float(str(x).replace(",", "").strip())
        return int(v) if v == int(v) else v
    except (TypeError, ValueError):
        return None


def _rows_to_lines(rows):
    """rows = list of cell lists; detect the header, qty column and description column."""
    rows = [[("" if c is None else str(c)).strip() for c in r] for r in rows if any(str(c or "").strip() for c in r)]
    if not rows:
        return []
    head, qi, di = None, None, None
    for i, r in enumerate(rows[:15]):
        low = [c.lower() for c in r]
        q = next((j for j, c in enumerate(low) if re.fullmatch(r"(qty|quantity|qnty|units?|count)\.?", c)), None)
        d = next((j for j, c in enumerate(low) if re.search(r"desc", c)), None)
        if d is None:  # a 'Product #' / part-number column is never the description
            d = next((j for j, c in enumerate(low) if re.search(r"product|item name|component|option name", c)
                      and not re.search(r"#|number|\bno\b|\bpn\b|sku|part", c)), None)
        if q is not None and d is not None:
            head, qi, di = i, q, d
            break
    out = []
    for r in rows[(head + 1 if head is not None else 0):]:
        if head is None:
            q = next((_num(c) for c in r if _num(c) is not None), 1)
            d = max(r, key=len)
        else:
            q = _num(r[qi]) if qi < len(r) else None
            d = r[di] if di < len(r) else ""
        if d and len(d) > 2:
            out.append(dict(qty=q if q is not None else 1, text=d, raw=" | ".join(c for c in r if c)))
    return out


_QX = re.compile(r"^\s*(\d+(?:\.\d+)?)\s*(?:x|X|×|pcs?|ea)?\s+(.+)$")


def load(path):
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        try:
            return _rows_to_lines(list(wb[wb.sheetnames[0]].iter_rows(values_only=True)))
        finally:
            wb.close()
    with open(path, encoding="utf-8-sig", errors="replace") as f:
        text = f.read()
    if ext == ".csv":
        return _rows_to_lines(list(csv.reader(text.splitlines())))
    out = []
    for raw in text.splitlines():
        s = raw.strip()
        if not s or s.startswith("#"):
            continue
        if "\t" in s:
            cells = [c.strip() for c in s.split("\t") if c.strip()]
            q = next((_num(c) for c in cells if _num(c) is not None), 1)
            d = " ".join(c for c in cells if _num(c) is None)
            out.append(dict(qty=q, text=d, raw=s))
            continue
        m = _QX.match(s)
        out.append(dict(qty=_num(m.group(1)), text=m.group(2).strip(), raw=s) if m else dict(qty=1, text=s, raw=s))
    return out


# ------------------------------------------------------------------ classification + spec extraction
CATS = [  # order matters: first match wins
    ("service", r"ProSupport|Care ?Pack|Foundation Care|Tech Care|Pointnext|SNTC|Smart Net|warranty|\bNBD\b|next business day|24x7|\b4[- ]?h(ou)?r\b|support services|installation service|deployment service"),
    ("software", r"Windows Server|\bCALs?\b|VMware|vSphere|vSAN|Nutanix (NCI|NCM|AOS|NUS)|Red Hat|RHEL|SUSE|licen[cs]e|subscription"),
    ("mgmt", r"iDRAC|\biLO\b|OneView|OpenManage|Intersight|IPMI|XClarity|\bBMC\b"),
    ("gpu", r"\bGPU\b|NVIDIA (L4|L40S?|A2|A16|A30|A40|H100|H200|RTX|Tesla)|\bL40S\b|RTX PRO|Instinct MI"),
    ("fc", r"Fib(re|er) Channel|\bFC\b.*(HBA|adapter)|\b(16|32|64)\s?G(b|FC)\b.*(HBA|Port)|QLogic|Emulex|LPe\d"),
    ("boot", r"\bBOSS\b|NS204|M\.2|boot (optimized|device|drive|controller)"),
    ("controller", r"PERC|\bRAID\b|\bHBA\b|MegaRAID|Smart ?Array|\bMR\d{3}|\bSR\d{3}\b.*Controller|storage controller|tri-?mode|\bH\d{3}\b"),
    ("optic", r"transceiver|\bSFP\+? ?(SR|LR)|\bSR4\b|\bLR4\b|\bDAC\b|twinax|direct attach|\bAOC\b|optic"),
    ("nic", r"Ethernet|\bNIC\b|\bOCP\b|\bLOM\b|\d+ ?Gb(E|ps)?\b.*(port|adapter)|ConnectX|Broadcom 5\d{4}|BCM5\d{3}|Intel (E8|X7|E6|I35)\d{2}|\bSFP28\b|\bQSFP"),
    ("cpu", r"Xeon|EPYC|\bprocessor\b|\bCPU\b"),
    ("memory", r"\bR?DIMM\b|DDR[45]|\bmemory\b|\bRAM\b"),
    ("drive", r"\b(SSD|HDD|NVMe)\b|hard drive|solid state|\b\d+(\.\d+)?\s?TB\b.*(SAS|SATA|U\.[23]|E3\.S)|\b(SAS|SATA)\b.*\b\d+(\.\d+)?\s?(TB|GB)\b|\b\d+[kK] ?RPM\b|7\.2K|10K|15K"),
    ("psu", r"power supply|\bPSU\b|\b\d{3,4}\s?W\b.*(Platinum|Titanium|hot.?plug|redundant|Flex Slot|CRPS)|Platinum|Titanium"),
    ("cord", r"power cord|line cord|jumper cord|C13|C14|C19|C20|NEMA"),
    ("rails", r"\brails?\b|\bCMA\b|cable management arm|bezel|rack mount kit"),
    ("platform", r"PowerEdge|ProLiant|\bDL\d{3}\b|\bML\d{3}\b|\bR\d{3,4}(xs|xd|xa)?\b|\bT\d{3}\b|UCS|\bC2[24]0\b|\bB200\b|\bX210c\b|Supermicro|\bSYS-|\bAS-\d|Nutanix NX|\bNX-\d|chassis|server\b|\bnode\b"),
]


def classify(t):
    for cat, rx in CATS:
        if re.search(rx, t, I):
            return cat
    return "other"


def _f(rx, t, cast=str, g=1):
    m = re.search(rx, t, I)
    return cast(m.group(g)) if m else None


def spec(cat, t):
    s = {}
    if cat == "cpu":
        s["vendor"] = "amd" if re.search(r"EPYC|AMD", t, I) else ("intel" if re.search(r"Xeon|Intel", t, I) else None)
        s["model"] = _f(r"(?:Xeon[^0-9]{0,24}|EPYC\s*)(\d{4}[A-Z+]{0,3})", t)
        s["cores"] = _f(r"(\d{1,3})\s*[- ]?\s*(?:C\b|cores?\b)", t, int)
        s["ghz"] = _f(r"(\d\.\d{1,2})\s*GHz", t, float)
        s["watts"] = _f(r"(\d{2,3})\s*W\b", t, int)
    elif cat == "memory":
        s["gb"] = _f(r"(\d{1,3})\s*GB", t, int)
        s["mts"] = _f(r"(\d{4})\s*(?:MT/s|MHz|MT)", t, int)
        s["kind"] = "udimm" if re.search(r"UDIMM|unbuffered", t, I) else "rdimm"
    elif cat in ("drive", "boot"):
        tb = _f(r"(\d+(?:\.\d+)?)\s*TB", t, float)
        gb = _f(r"(\d{3,4})\s*GB", t, int)
        s["tb"] = tb if tb else (round(gb / 1000, 2) if gb else None)
        s["iface"] = "nvme" if re.search(r"NVMe|U\.[23]|E3\.S|PCIe", t, I) else ("sas" if re.search(r"\bSAS\b", t, I) else ("sata" if re.search(r"SATA", t, I) else None))
        s["form"] = "3.5" if re.search(r"3\.5|LFF", t, I) else ("2.5" if re.search(r"2\.5|SFF|U\.[23]", t, I) else ("e3s" if re.search(r"E3\.S", t, I) else None))
        s["u3"] = bool(re.search(r"U\.3", t))
        s["media"] = "hdd" if re.search(r"HDD|RPM|7\.2K|10K|15K|hard drive", t, I) else "ssd"
        s["endurance"] = "mu" if re.search(r"mixed use|\bMU\b|3 ?DWPD", t, I) else ("ri" if re.search(r"read intensive|\bRI\b|1 ?DWPD", t, I) else None)
        if cat == "boot":  # 'BOSS-N1 ... with 2 M.2 480GB' = a kit plus two drives
            s["drives"] = _f(r"(?<![\w.])(\d)\s*x?\s*(?:M\.2|NVMe|SSD|drives?)", t, int) or (2 if re.search(r"RAID ?1|mirror", t, I) else None)
            if re.search(r"NS204", t, I):  # HPE NS204i-u = hot-plug RAID 1 pair of 480GB NVMe M.2
                s.update(drives=2, iface="nvme", tb=s.get("tb") or 0.48)
    elif cat == "controller":
        s["ports"] = _f(r"(\d{1,2})\s*i\b", t, int) or _f(r"(\d{1,2})[- ]port", t, int)
        s["cache_gb"] = _f(r"(\d)\s*GB", t, int)
        s["hba"] = bool(re.search(r"\bHBA\b|H3[45]5|pass.?through|\bIT\b", t, I))
        s["trimode"] = bool(re.search(r"tri-?mode|MR416|H965|H755N", t, I))
    elif cat in ("nic", "optic"):
        sp = re.findall(r"(\d{1,3})\s*(?:/\s*(\d{1,3}))?\s*G(?:b|bE|BASE|igabit)", t, I)
        s["speeds"] = sorted({int(x) for a in sp for x in a if x})
        s["ports"] = _f(r"(\d)\s*[- ]?(?:x\s*)?port", t, int) or (2 if re.search(r"dual", t, I) else 4 if re.search(r"quad", t, I) else None)
        s["media"] = "baset" if re.search(r"BASE-?T|RJ-?45|copper", t, I) else ("qsfp" if re.search(r"QSFP", t, I) else ("sfp" if re.search(r"SFP|optical|fib(re|er)", t, I) else None))
        s["ocp"] = bool(re.search(r"\bOCP ?(3(\.0)?)?\b|\bLOM\b|mezz|FlexLOM", t, I))
        s["vendor"] = next((v for v, rx in (("broadcom", r"Broadcom|BCM"), ("intel", r"\bIntel\b"), ("nvidia", r"Mellanox|ConnectX|NVIDIA")) if re.search(rx, t, I)), None)
    elif cat == "psu":
        s["watts"] = _f(r"(\d{3,4})\s*W", t, int)
        s["titanium"] = bool(re.search(r"Titanium", t, I))
    elif cat == "gpu":
        s["model"] = _f(r"(L4|L40S|L40|A16|A2|A30|A40|H100|H200|RTX PRO \d{4}|RTX \d{4}|MI\d{3})", t)
        s["vram"] = _f(r"(\d{2,3})\s*GB", t, int)
    elif cat == "fc":
        s["gbps"] = _f(r"(16|32|64)\s?G", t, int)
        s["ports"] = _f(r"(\d)\s*[- ]?port", t, int) or (2 if re.search(r"dual", t, I) else None)
    return {k: v for k, v in s.items() if v not in (None, [], "")}


# ------------------------------------------------------------------ platform choice
BASE_MTM = {"SR630 V4": "7DG9CTO1WW", "SR650 V4": "7DGDCTO1WW", "SR650a V4": "7DGDCTO2WW", "SR635 V3": "7D9GCTO1WW",
            "SR645 V3": "7D9CCTO1WW", "SR655 V3": "7D9ECTO1WW", "SR665 V3": "7D9ACTO1WW", "SR250 V3": "7DCLCTO1WW",
            "ST250 V3": "7DCECTO1WW", "ST50 V3": "7DF3CTO1WW", "HX630 V4": "7DG3CTO1WW", "HX650 V4": "7DG4CTO1WW",
            "HX650 V4 Storage": "7DG4CTO2WW", "FX630 V4": "7DPLCTO1WW"}


def form_factor(t):
    t2 = t.upper()
    if re.search(r"\bR(2|3|4|6)\d{2}\b|\bR6\d{3}\b|\bDL3[26]0\b|\bC220\b|\b1U\b|SYS-1\d", t2): return "1U"
    if re.search(r"\bR7\d{2}\b|\bR7\d{3}\b|\bDL3[89]0\b|\bDL385\b|\bC240\b|\b2U\b|SYS-[25]\d|NX-[38]\d{3}", t2): return "2U"
    if re.search(r"\bT\d{3}\b|\bML\d{2,3}\b|tower|MicroServer", t, I): return "tower"
    if re.search(r"\bB200\b|\bX210c\b|blade", t, I): return "blade"
    return None


def guess_platforms(lines, workload=None):
    plat = next((l for l in lines if l["cat"] == "platform"), None)
    cpus = [l for l in lines if l["cat"] == "cpu"]
    sockets = int(cpus[0]["qty"]) if cpus and cpus[0]["qty"] in (1, 2) else (2 if len(cpus) >= 2 else 1 if cpus else None)
    vendor = next((l["spec"].get("vendor") for l in cpus if l["spec"].get("vendor")), None)
    ff = form_factor(plat["text"]) if plat else None
    drives35 = any(l["cat"] == "drive" and l["spec"].get("form") == "3.5" for l in lines)
    gpus = sum(l["qty"] for l in lines if l["cat"] == "gpu")
    why, out = [], []
    if workload == "nutanix" or (plat and re.search(r"Nutanix|\bNX-\d", plat["text"], I)):
        out = ["HX650 V4", "HX630 V4"] if ff != "1U" else ["HX630 V4", "HX650 V4"]
        if drives35: out.insert(0, "HX650 V4 Storage")
        why.append("Nutanix workload: ThinkAgile HX (single-socket source nodes go to HX650 V4; HX630 V4 needs 2 CPUs)")
    elif ff == "tower":
        out = ["ST250 V3", "ST50 V3"]
        why.append("tower source: ST250 V3 (8C max, SAS/SATA, hot-swap) or ST50 V3 (smallest)")
    elif gpus >= 2:
        out = ["SR650a V4", "SR655 V3", "SR665 V3"]
        why.append(f"{gpus:g} GPUs: SR650a V4 (Intel, Server Edition GPUs) or the AMD 'for AI' SR655/665 V3 models")
    elif vendor == "amd":
        out = (["SR635 V3", "SR645 V3"] if sockets == 1 else ["SR645 V3"]) if ff == "1U" else (["SR655 V3", "SR665 V3"] if sockets == 1 else ["SR665 V3"])
        why.append(f"AMD source, {sockets or '?'} socket(s), {ff or 'form factor unknown'}")
    else:
        out = ["SR630 V4", "SR650 V4"] if ff == "1U" else ["SR650 V4", "SR630 V4"]
        if drives35 and "SR650 V4" in out:
            out.remove("SR650 V4"); out.insert(0, "SR650 V4")
            why.append("3.5-inch drives: SR650 V4 (the 1U holds only 4x 3.5in)")
        why.append(f"Intel (or unknown) source, {ff or 'form factor unknown'}: current Xeon 6 platforms")
    comp = compete_hint(plat["text"]) if plat else []
    return out, why, comp, plat


_MODEL = re.compile(r"\b(R\d{3,4}(?:xs|xd|xa|xd2)?|T\d{3}|DL\d{3}(?:\s*Gen\s*\d{1,2})?|ML\d{2,3}(?:\s*Gen\s*\d{1,2})?|"
                    r"C2[24]0\s*M\d|B200\s*M\d|X210c\s*M\d|SYS-[\w-]+|AS\s*-\s*[\w-]+|NX-\d{4}\w*)\b", I)
_VENDOR = (("dell", r"Dell|PowerEdge"), ("hpe", r"\bHPE?\b|ProLiant"), ("cisco", r"Cisco|UCS"),
           ("supermicro", r"Supermicro|\bSYS-|\bAS\s*-"), ("nutanix", r"Nutanix|\bNX-\d"))


def compete_hint(text):
    """Official Lenovo competitor map suggestions via kb/compete.py (live, cached); [] when unavailable."""
    try:
        import compete
    except ImportError:
        return []
    m = _MODEL.search(text)
    if not m:
        return []
    ven = next((v for v, rx in _VENDOR if re.search(rx, text, I)), "")
    return compete.suggest((ven + " " + m.group(1)).strip())


# ------------------------------------------------------------------ candidate search
SEC = {  # DCSC section-name patterns per category
    "cpu": r"^PROCESSOR$", "memory": r"^MEMORY$", "drive": r"NORAID_HDD|NVME_DRIVES|^HDD$|DRIVES$", "boot": r"^M\.?2_SSD|M2_SSD|7MM",
    "controller": r"RAID_CONTROLLER|NON_RAID_CONTROLLER|HBA", "nic": r"OCP|NETWORK_ADAPTER", "optic": r"TRANSCEIVER|EXTERNAL_CABLES",
    "psu": r"POWERSUPPLY", "gpu": r"GPU_ADAPTER", "fc": r"FIBER_CHANNEL", "cord": r"LINECORD|RACK_POWERCABLE",
}


def _ok(d, rx):
    return re.search(rx, d, I) is not None


def score(cat, s, d):
    """Higher = closer. Hard mismatches return None."""
    sc = 0
    if cat == "cpu":
        if s.get("vendor") == "amd" and not _ok(d, r"EPYC"): return None
        if s.get("vendor") == "intel" and not _ok(d, r"Xeon"): return None
        if s.get("model") and s["model"].upper() in d.upper(): sc += 20
        c = _f(r"(\d{1,3})C\b", d, int)
        if s.get("cores") and c:
            sc += 10 if c == s["cores"] else max(0, 6 - abs(c - s["cores"]) // 4)
        g = _f(r"(\d\.\d+)GHz", d, float)
        if s.get("ghz") and g: sc += max(0, 4 - abs(g - s["ghz"]) * 4)
    elif cat == "memory":
        g = _f(r"(\d{1,3})GB", d, int)
        if s.get("gb") and g != s["gb"]: return None
        sc += 10
        if s.get("kind") == "udimm" and not _ok(d, r"UDIMM"): sc -= 5
        if s.get("mts") and str(s["mts"]) in d: sc += 3
        if _ok(d, r"MRDIMM|3DS"): sc -= 4
    elif cat in ("drive", "boot"):
        tb = _f(r"(\d+(?:\.\d+)?)TB", d, float) or ((_f(r"(\d{3,4})GB", d, int) or 0) / 1000)
        if s.get("tb"):
            if not tb: return None
            r = tb / s["tb"]
            if r < 0.95 or r > 2.1: return None
            sc += 10 if 0.95 <= r <= 1.06 else 4
        if s.get("iface") == "nvme" and not _ok(d, r"NVMe"): return None
        if s.get("iface") in ("sas", "sata") and _ok(d, r"NVMe"): return None
        if s.get("iface") == "sas" and _ok(d, r"\bSAS\b"): sc += 3
        if s.get("iface") == "sata" and _ok(d, r"SATA"): sc += 3
        if s.get("media") == "hdd" and not _ok(d, r"HDD"): return None
        if s.get("media") == "ssd" and _ok(d, r"HDD"): return None
        if s.get("form") in ("2.5", "3.5"):
            other = "3.5" if s["form"] == "2.5" else "2.5"
            if s["form"] + '"' in d: sc += 2
            elif other + '"' in d: sc -= 6  # wrong form factor for the source chassis
        if s.get("endurance") == "mu" and _ok(d, r"Mixed Use"): sc += 3
        if s.get("endurance") == "ri" and _ok(d, r"Read Intensive"): sc += 3
        if _ok(d, r"U\.3") and not s.get("u3"): sc -= 3  # U.3 kills TCE and is rarely required
        if _ok(d, r"\bSED\b|FIPS"): sc -= 2
    elif cat == "controller":
        if s.get("hba"):
            sc += 6 if _ok(d, r"HBA|440") else 0
        else:
            sc += 6 if _ok(d, r"RAID 9\d0") else 2 if _ok(d, r"RAID") else 0
        p = _f(r"(\d{1,2})i\b", d, int)
        if s.get("ports") and p: sc += 5 if p == s["ports"] else 1
        if s.get("cache_gb") and _ok(d, rf"{s['cache_gb']}GB"): sc += 3
        if _ok(d, r"Tri-?Mode|U\.3") and not s.get("trimode"): sc -= 4
    elif cat in ("nic", "optic"):
        sp = [int(x) for x in re.findall(r"(\d{1,3})G(?:bE|b|BASE)", d, I)]
        if s.get("speeds"):
            if not set(sp) & set(s["speeds"]) and not (max(s["speeds"]) == 10 and 25 in sp and s.get("media") != "baset"): return None
            sc += 6
        if s.get("media") == "baset" and not _ok(d, r"BASE-?T"): return None
        if s.get("media") in ("sfp", "qsfp") and _ok(d, r"BASE-?T"): return None
        if s.get("media") == "qsfp" and _ok(d, r"QSFP"): sc += 3
        p = _f(r"(\d)-?[Pp]ort", d, int)
        if s.get("ports") and p: sc += 4 if p == s["ports"] else 0
        if cat == "nic" and s.get("ocp") is not None:
            sc += 3 if ("OCP" in d) == s["ocp"] else 0
        if s.get("vendor") == "broadcom" and _ok(d, r"Broadcom"): sc += 2
        if s.get("vendor") == "intel" and _ok(d, r"Intel"): sc += 2
        if s.get("vendor") == "nvidia" and _ok(d, r"Mellanox|NVIDIA|ConnectX"): sc += 2
    elif cat == "psu":
        w = _f(r"(\d{3,4})W", d, int)
        if s.get("watts") and w:
            if w < s["watts"] * 0.6: return None
            sc += max(0, 10 - abs(w - s["watts"]) / 100)
        if _ok(d, r"230V/115V"): sc += 2
        if s.get("titanium") and _ok(d, r"Titanium"): sc += 1
    elif cat == "gpu":
        if s.get("model") and s["model"].replace(" ", "").upper() not in d.replace(" ", "").upper(): return None
        sc += 10
    elif cat == "fc":
        if s.get("gbps") and not _ok(d, rf"{s['gbps']}Gb"): return None
        sc += 5
    return sc


def candidates(db, mtm, cat, s, tce_first=True, n=3):
    rx = SEC.get(cat)
    if not rx:
        return []
    seen, out = set(), []
    for tab, sub, sec, c, d, mn, mx, tc, wd, ty in dcsc_rules.options(db, mtm):
        if not re.search(rx, sec, I) or c in seen or wd:
            continue
        if cat == "drive" and _ok(d, r"M\.2|7mm"): continue
        if cat == "boot" and not _ok(d, r"M\.2|7mm"): continue
        if cat == "nic" and not _ok(d, r"Ethernet|Adapter"): continue
        sc = score(cat, s, d)
        if sc is None:
            continue
        seen.add(c)
        out.append(dict(fc=c, descr=d, tce=bool(tc), section=sec, score=round(sc + (3 if (tc and tce_first) else 0), 1)))
    out.sort(key=lambda x: (-x["score"], not x["tce"]))  # ties go to the TCE part
    return out[:n]


# ------------------------------------------------------------------ companions and notes
def companions(lines, picks):
    notes = []
    cats = {l["cat"] for l in lines}
    sfp = sum(l["qty"] * (l["spec"].get("ports") or 2) for l in lines if l["cat"] == "nic" and l["spec"].get("media") in ("sfp", "qsfp", None) and l["spec"].get("media") != "baset")
    if sfp and "optic" not in cats:
        notes.append(f"No optics or DACs on the source BOM: ~{sfp:g} SFP/QSFP ports per server need a transceiver or DAC (Lenovo BOMs ship none). Ask fiber vs DAC and the distance.")
    if "boot" not in cats:
        notes.append("No boot device on the source BOM: most Lenovo builds add a mirrored M.2 pair (V4 Intel: CCCZ rear hot-swap kit + 2 drives; V3 AMD: B8P9 + 2x 480GB).")
    if any(p and re.search(r"RAID 9\d0", p["descr"], I) for p in picks):
        notes.append("A 940-series RAID adapter pulls in the SuperCap (AUNP) and its cable; check DCSC added them.")
    if "gpu" in cats:
        notes.append("GPUs force companions: performance fans, DIMM fillers in empty slots, GPU power cables and riser cages, and often a 230V-only PSU. Check the GPU facts (kb.py fact GPU).")
    if "mgmt" in cats:
        notes.append("iDRAC Enterprise / iLO Advanced / Intersight map to XCC Platinum (V3, SBCV) or XCC3 Premier (V4, SCY0), plus XClarity One for fleet management.")
    if "service" in cats:
        notes.append("Map services by response, not by name: Dell ProSupport NBD / HPE NBD -> Premier NBD; 4-hour -> Premier 24x7 4hr; add Keep Your Drive where drives hold regulated data.")
    return notes


def related_facts(lines, limit=5):
    try:
        from kb_facts import FACTS
    except Exception:
        return []
    toks = set()
    for l in lines:
        for t in re.findall(r"\b(R\d{3,4}\w*|DL\d{3}|C2[24]0|B200|SYS-\w+|NX-\d+|PERC|H\d{3}|BOSS|NS204|MR\d{3}|E810|X710|57\d{3}|ConnectX-\d|L4|L40S|RTX PRO|6\d{3}P|9\d{3}P?)\b", l["text"], I):
            toks.add(t.lower())
    scored = []
    for f in FACTS:
        blob = (f[1] + " " + f[3] + " " + f[4]).lower()
        hit = sum(1 for t in toks if t in blob)
        if hit:
            scored.append((hit, f[0], f[3]))
    scored.sort(key=lambda x: -x[0])
    return [(t, ti) for _, t, ti in scored[:limit]]


# ------------------------------------------------------------------ main flow
def run(path, platform=None, mtm=None, nodes=None, tce=False, workload=None):
    db = dcsc_rules.load()
    lines = load(path)
    for l in lines:
        l["cat"] = classify(l["text"])
        l["spec"] = spec(l["cat"], l["text"])
    notes = []
    plat_line = next((l for l in lines if l["cat"] == "platform"), None)
    n = nodes or (int(plat_line["qty"]) if plat_line and plat_line["qty"] and plat_line["qty"] > 1 else None)
    if n and n > 1 and not nodes:
        notes.append(f"The platform line has qty {n}: treating the file as an ORDER TOTAL for {n} servers and dividing "
                     f"quantities by {n} (HPE/Dell exports list totals). Pass --nodes 1 if the quantities are already per server.")
    if n and n > 1:
        for l in lines:
            if l["cat"] in ("service", "software"):
                continue
            q = l["qty"] / n
            if q != int(q):
                notes.append(f"'{l['text'][:50]}' qty {l['qty']:g} does not divide evenly by {n}.")
            l["qty"] = int(q) if q == int(q) else round(q, 2)
    choices, why, comp, _ = guess_platforms(lines, workload)
    target = None
    if mtm:
        target = mtm.upper()
    elif platform:
        hits = dcsc_rules.resolve(db, platform)
        target = BASE_MTM.get(platform.upper().replace("THINKSYSTEM ", "").replace("THINKAGILE ", ""), None) or (hits[0] if len(hits) == 1 else None)
        if not target:
            k = next((k for k in BASE_MTM if k.lower() == platform.lower().strip()), None)
            target = BASE_MTM.get(k) if k else None
    else:
        target = BASE_MTM.get(choices[0]) if choices else None
    if not target or target not in db["models"]:
        sys.exit(f"could not resolve a Lenovo MTM (platform={platform!r}, mtm={mtm!r}); pass --mtm, e.g. --mtm 7DG9CTO1WW")
    rows, picks = [], []
    for l in lines:
        cands = candidates(db, target, l["cat"], l["spec"], tce_first=True)
        rows.append(dict(qty=l["qty"], cat=l["cat"], text=l["text"], spec=l["spec"], candidates=cands))
        picks.append(cands[0] if cands else None)
    bom = []
    for r, p in zip(rows, picks):
        if p:
            q = r["qty"] * r["spec"]["drives"] if r["cat"] == "boot" and r["spec"].get("drives") and r["qty"] == 1 else r["qty"]
            if q != r["qty"]:
                r["qty"] = q
                notes.append(f"Boot line read as {q:g} drives ('{r['text'][:40]}'): quote {q:g}x {p['fc']} plus the M.2 enablement/RAID kit for the platform.")
            bom.append(dict(fc=p["fc"], descr=p["descr"], qty=q, tce=p["tce"]))
    findings = []
    try:
        import rules
        findings = rules.validate(bom, mtm=target, nodes=n or 1, tce=tce or None, workload=workload)
    except Exception as e:
        findings = [dict(rule="validator-unavailable", severity="info", message=f"rules.py not run: {e}", evidence=[])]
    non_tce = [p["fc"] for p in picks if p and not p["tce"]]
    return dict(target=target, target_name=db["models"][target]["name"], snapshot=db["meta"]["crawled"], platform_choices=choices,
                platform_why=why, compete=comp, nodes=n or 1, notes=notes, rows=rows, first_pick_bom=bom,
                non_tce_first_picks=non_tce, companions=companions(lines, picks), findings=findings,
                related_facts=related_facts(lines))


def show(r):
    print(f"# Competitor BOM -> Lenovo worksheet\n")
    print(f"Target: {r['target']} {r['target_name']}  (DCSC rules snapshot {r['snapshot']}; servers: {r['nodes']})")
    print(f"Platform candidates: {', '.join(r['platform_choices'])}")
    for w in r["platform_why"]: print(f"  - {w}")
    if r["compete"]:
        print("Official Lenovo competitor map:")
        for c in r["compete"][:5]: print(f"  - {c}")
    for n in r["notes"]: print(f"NOTE: {n}")
    print("\n| # | qty | source line | category | proposed Lenovo FC (TCE) |\n|---|---|---|---|---|")
    for i, row in enumerate(r["rows"], 1):
        c = row["candidates"]
        prop = "<br>".join(f"{x['fc']} {'TCE' if x['tce'] else 'not TCE'}: {x['descr'][:60]}" for x in c) if c else ("map by hand" if row["cat"] not in ("service", "software", "mgmt", "rails", "platform") else "see notes")
        print(f"| {i} | {row['qty']:g} | {row['text'][:70].replace('|', '/')} | {row['cat']} | {prop} |")
    if r["non_tce_first_picks"]:
        print(f"\nFirst picks that are NOT TCE in the snapshot: {', '.join(r['non_tce_first_picks'])} (TCE is all-or-nothing).")
    print("\n## Forced companions and gaps")
    for c in r["companions"]: print(f"- {c}")
    print("\n## Validator findings on the first-pick BOM")
    for f in r["findings"]:
        print(f"- [{f['severity']}] {f['rule']}: {f['message']}")
    if r["related_facts"]:
        print("\n## Related verified facts (python kb.py fact <topic words>)")
        for t, ti in r["related_facts"]: print(f"- {t}: {ti}")
    print("\nEverything above is a proposal from a dated snapshot. Build it in DCSC and check the export for BU1E.")


def main(argv):
    for s in (sys.stdout, sys.stderr):
        try: s.reconfigure(encoding="utf-8", errors="replace")
        except Exception: pass
    a = list(argv)
    if a and a[0] == "match":
        a = a[1:]
    if not a or a[0] in ("-h", "--help"):
        print(__doc__); return 0
    def opt(k, cast=str):
        if k in a:
            i = a.index(k); v = a[i + 1]; del a[i:i + 2]; return cast(v)
        return None
    def flag(k):
        if k in a:
            a.remove(k); return True
        return False
    kw = dict(platform=opt("--platform"), mtm=opt("--mtm"), nodes=opt("--nodes", int), workload=opt("--workload"), tce=flag("--tce"))
    js = flag("--json")
    if len(a) != 1:
        print(__doc__); return 2
    r = run(a[0], **kw)
    if js: print(json.dumps(r, indent=1, default=str))
    else: show(r)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
