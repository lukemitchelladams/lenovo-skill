"""
Deterministic Lenovo BOM validators. Pure functions over a bill of materials:
no network, no model, no catalog. A rule FLAGS, it never edits the BOM.

usage:
  python rules.py check <bom-file> [--mtm X] [--nodes N] [--workload W] [--tce] [--raid N]
                                   [--cluster-totals] [--allow-2node] [--json]
  python rules.py diff <intended> <resolved> [--nodes N]   intended vs what DCSC resolved
  python rules.py list                                     every rule id + severity

BOM file: a DCSC export .xlsx (Quote sheet: A feature code, C description, E qty; price ignored)
or a text file, one line each:   2 x BN2T ThinkSystem Broadcom 57414 ...   or   BN2T<TAB>desc<TAB>2
Text files may start with directives:  # mtm: 7DG4CTO1WW   # nodes: 3   # workload: nutanix
                                       # tce: yes   # raid: 5   # cluster_totals: yes
Quantities are PER NODE. Exports carry cluster totals, so they are divided by the node count
(the qty on the MTM line) on load. Text files are taken as per-node unless --cluster-totals.

Severity: block = will not validate or is wrong; warn = check before quoting; info = context.
Exit code 1 if any block. Platform facts come from data/platform-specs.json (specs.py).
"""
import collections, json, os, re, sys

HERE = os.path.dirname(os.path.abspath(__file__))
if HERE not in sys.path:
    sys.path.insert(0, HERE)
import specs

WORKLOADS = ("nutanix", "vsan", "s2d", "windows")   # only 'nutanix' changes rule behaviour today
MTM_RE = re.compile(r"^7[A-Z0-9]{3}CTO[0-9A-Z]{1,2}WW$")
I = re.I

# ---------------------------------------------------------------- hand-verified TCE verdicts
# Beat any inferred evidence: the Top Choice Express list rotates, so a verdict is only good
# for its date. basis: panel = seen in the live DCSC panel, doc = Lenovo documentation,
# rep = confirmed by a Lenovo rep. mtm "*" = every platform.
TCE_OVERRIDES = [
    dict(fc="C5QQ", mtm="7DG4CTO1WW", tce=False, verified_on="2026-08-24", basis="panel",
         note="Xeon 6505P 12C not selectable under the TCE filter on HX650 V4; the list rotated. Smallest single-socket TCE CPU then was 16C (C5RD / C5QV)."),
    dict(fc="C5R6", mtm="7DG3CTO1WW", tce=True, verified_on="2026-08-19", basis="panel",
         note="Xeon 6507P 8C is TCE on HX630 V4 (older static snapshots said otherwise)."),
    dict(fc="C528", mtm="7DCXCTO1WW", tce=False, verified_on="2026-08-20", basis="panel",
         note="32GB UDIMM looked TCE-proven on ST250 V3 from earlier exports, but the live memory panel offered only the 16GB part. Export evidence is not proof it is selectable in every config."),
    dict(fc="C1X9", mtm="*", tce=False, verified_on="2026-08-06", basis="panel",
         note="7.68TB RI NVMe is not TCE on V4 HX. 15.36TB/node is built as 4x C3Q9 3.84TB, never 2x C1X9."),
    dict(fc="CCVY", mtm="*", tce=False, verified_on="2026-08-13", basis="doc", note="6.4TB NVMe not TCE."),
    dict(fc="CCW0", mtm="*", tce=False, verified_on="2026-08-13", basis="doc",
         note="1.6TB Mixed Use NVMe not TCE. CCVZ 3.2TB MU is TCE and swaps 2:1 at identical capacity."),
    dict(fc="BZ7D", mtm="*", tce=False, verified_on="2026-08-13", basis="panel",
         note="96GB DIMM is not TCE on any platform, and does not build on HX650 V4 Storage at all."),
    dict(fc="C0TY", mtm="*", tce=False, verified_on="2026-08-13", basis="doc", note="32GB MRDIMM not TCE."),
    dict(fc="C0TX", mtm="*", tce=False, verified_on="2026-08-13", basis="doc", note="64GB MRDIMM not TCE."),
    dict(fc="BS2C", mtm="*", tce=False, verified_on="2026-08-12", basis="panel",
         note="L4 24GB GPU not TCE. With the TCE filter on, the GPU category shows only BP4X."),
    dict(fc="C520", mtm="*", tce=True, verified_on="2026-08-19", basis="panel",
         note="Xeon 6325P 4C 55W is TCE (older static snapshots said otherwise)."),
]

def tce_verdict(fc, mtm=""):
    """True/False from TCE_OVERRIDES (exact MTM beats '*'), None = no hand-verified verdict."""
    fc, mtm = (fc or "").upper(), (mtm or "").upper()
    hits = [o for o in TCE_OVERRIDES if o["fc"] == fc and o["mtm"] in (mtm, "*")]
    hits.sort(key=lambda o: o["mtm"] == "*")
    return hits[0]["tce"] if hits else None

# ---------------------------------------------------------------- BOM lines
class Bom(list):
    """list of {fc, descr, qty} dicts (per-node qty) plus .meta detected from the file."""
    def __init__(self, it=(), meta=None):
        super().__init__(it)
        self.meta = meta or {}

def _num(x):
    try:
        v = float(str(x).replace(",", "").strip())
    except ValueError:
        return None
    return int(v) if v == int(v) else v

def _fq(q):
    return str(int(q)) if float(q) == int(q) else f"{q:g}"

SVC_FC = re.compile(r"^(5WS7|5PS7|5MS7|5AS7|7Q0\d)", I)
SVC_DESC = re.compile(r"premier|warranty|\bKYD\b|\bresp\b|deployment|\bS&S\b|committed service repair|hw installation", I)
# hardware that merely mentions its own warranty ("... Rail Kit, Lifetime Warranty") is still hardware
HW_DESC = re.compile(r"indicator|label|bracket|cable|filler|tray|shroud|heatsink|riser|backplane|cage|rail", I)

def _service_like(fc, d):
    return bool(SVC_FC.match(fc or "")) or (bool(SVC_DESC.search(d)) and not HW_DESC.search(d))

def section_of(fc, d):
    """Best-effort DCSC-style section from the description (exports carry no section column)."""
    if _service_like(fc, d): return "Services"
    if re.search(r"Power Supply|Power Cord|Line Cord|Jumper Cord|Rack Power Cable|\bPDU\b", d, I): return "Power"
    if re.search(r"M\.2", d, I): return "Storage"
    if re.search(r"Heatsink|(Xeon|EPYC).*Processor|Processor$", d, I): return "Processor"
    if re.search(r"\b\w*DIMM\b|TruDDR", d, I) and not re.search(r"Filler|Blank|Cable|Riser", d, I): return "Memory"
    if re.search(r"SSD|HDD|Backplane|\bRAID\b|\bHBA\b|Enablement|NVMe|\bSAS\b|\bSATA\b|Boot", d, I): return "Storage"
    if re.search(r"Windows Server|VMware|vSphere|Nutanix|License|XClarity|Red Hat|SUSE|Hyper-V", d, I): return "OS & software"
    if re.search(r"Ethernet|Adapter|Riser|OCP|Transceiver|SFP|GPU|InfiniBand|Fibre|Cable Kit|PCIe|Cage", d, I): return "PCI"
    return "Others"

def norm_line(x):
    """Accept {fc, descr, qty} or the long-key shape {code, description, qty, section, tc}."""
    fc = str(x.get("fc") or x.get("code") or "").strip().upper()
    d = str(x.get("descr") or x.get("description") or "").strip()
    q = _num(x.get("qty", 1))
    tce = x.get("tce", x.get("tc"))
    return dict(fc=fc, descr=d, qty=1 if q is None else q, tce=tce if isinstance(tce, bool) else None,
                section=x.get("section") or section_of(fc, d))

def is_service(c):
    return c["section"] == "Services" or _service_like(c["fc"], c["descr"])

# ---------------------------------------------------------------- loaders
def _finish(lines, meta):
    """Detect MTM + nodes from the MTM line; divide cluster totals down to per-node."""
    mtms = [l["fc"] for l in lines if MTM_RE.match(l["fc"])]
    if mtms and not meta.get("mtm"):
        meta["mtm"] = collections.Counter(mtms).most_common(1)[0][0]
    notes = meta.setdefault("notes", [])
    if len(set(mtms)) > 1:
        notes.append(f"{len(set(mtms))} different MTMs in the BOM ({', '.join(sorted(set(mtms)))}); validated as ONE node type, split the file per node type for a clean result.")
    mline = next((l for l in lines if l["fc"] == meta.get("mtm")), None)
    if mline and not meta.get("nodes") and mline["qty"] >= 1:
        meta["nodes"] = int(mline["qty"])
    n = meta.get("nodes")
    if meta.get("cluster_totals") and n and n > 1:
        for l in lines:
            if is_service(l): continue
            q = l["qty"] / n
            if q != int(q): notes.append(f"{l['fc']} qty {_fq(l['qty'])} does not divide evenly by {n} nodes.")
            l["qty"] = int(q) if q == int(q) else round(q, 3)
        notes.append(f"Quantities divided by {n} nodes (cluster totals to per-node).")
    elif not meta.get("cluster_totals") and mline and mline["qty"] > 1:
        notes.append(f"MTM line qty is {_fq(mline['qty'])}; quantities are treated as PER-NODE. Use --cluster-totals if they are cluster totals.")
    return Bom(lines, meta)

def load_xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb["Quote"] if "Quote" in wb.sheetnames else wb[wb.sheetnames[0]]
        lines = []
        for row in ws.iter_rows(values_only=True):
            c = ["" if v is None else str(v).strip() for v in row]
            if len(c) < 5 or not c[0] or not c[2]: continue
            q = _num(c[4])
            if q is None or q <= 0: continue
            lines.append(norm_line(dict(fc=c[0], descr=c[2], qty=q)))
    finally:
        wb.close()
    return _finish(lines, dict(source=os.path.basename(path), cluster_totals=True))

_DIRECTIVE = re.compile(r"^#\s*(mtm|nodes|workload|tce|raid|family|cluster_totals)\s*[:=]\s*(.+?)\s*$", I)
_QTY_X = re.compile(r"^(\d+(?:\.\d+)?)\s*[xX×]\s*([A-Za-z0-9][A-Za-z0-9_-]*)\s+(.+?)\s*$")
_BARE = re.compile(r"^([A-Z0-9][A-Z0-9_-]{3,11})\s*[-:]?\s+(\S.*?)\s*$")
_YES = ("1", "y", "yes", "true", "on")

def parse_text(text):
    lines, meta, skipped = [], {}, []
    for raw in text.splitlines():
        s = raw.strip()
        if not s: continue
        if s.startswith("#"):
            m = _DIRECTIVE.match(s)
            if m:
                k, v = m.group(1).lower(), m.group(2)
                meta[k] = (v.lower() in _YES) if k in ("tce", "cluster_totals") else (_num(v) if k in ("nodes", "raid") else v)
            continue
        f = [x.strip() for x in raw.split("\t")] if "\t" in raw else None
        rec = None
        if f and len(f) >= 3 and _num(f[-1]) is not None and f[0]:
            rec = (f[0], " ".join(f[1:-1]), _num(f[-1]))
        elif f and len(f) >= 3 and _num(f[0]) is not None:
            rec = (f[1], " ".join(f[2:]), _num(f[0]))
        elif f and len(f) == 2 and f[0]:
            rec = (f[0], f[1], 1)
        else:
            m = _QTY_X.match(s)
            if m:
                rec = (m.group(2), m.group(3), _num(m.group(1)))
            else:
                m = _BARE.match(s)
                tok = m.group(1) if m else ""
                if m and (any(ch.isdigit() for ch in tok) or len(tok) == 4) and tok.upper() == tok:
                    rec = (tok, m.group(2), 1)
        if rec and re.match(r"^[A-Za-z0-9][A-Za-z0-9_-]*$", rec[0]):
            lines.append(norm_line(dict(fc=rec[0], descr=rec[1], qty=rec[2])))
        else:
            skipped.append(s)
    if skipped:
        meta.setdefault("notes", []).append(f"{len(skipped)} line(s) not parsed and ignored, e.g. \"{skipped[0][:60]}\".")
    return _finish(lines, meta)

def load_bom(path):
    """DCSC export .xlsx or pasted-text file -> Bom (list of {fc, descr, qty, section, tce})."""
    if os.path.splitext(path)[1].lower() in (".xlsx", ".xlsm"):
        return load_xlsx(path)
    with open(path, encoding="utf-8-sig", errors="replace") as fh:
        b = parse_text(fh.read())
    b.meta["source"] = os.path.basename(path)
    return b

# ---------------------------------------------------------------- context + line helpers
def _is_hx(family, mtm):
    return bool(re.search(r"\bHX\d{3,4}", family or "", I) or re.match(r"7DG[34]|7D6|7DP", mtm or "", I))

def make_ctx(bom, mtm=None, family=None, nodes=None, workload=None, tce=None, raid=None, allow_2node=False):
    meta = getattr(bom, "meta", {}) or {}
    lines = [l if "section" in l and "descr" in l else norm_line(l) for l in bom]
    mtm = str(mtm or meta.get("mtm") or "").upper()
    fam_in = family or meta.get("family") or ""
    plat = specs.resolve(mtm=mtm or None, name=fam_in or None)
    fam = fam_in or (plat["name"] if plat else "")
    if tce is None: tce = meta.get("tce")
    if tce is None: tce = any(l["fc"] == "BU1E" for l in lines)
    n = nodes or meta.get("nodes")
    wl = str(workload or meta.get("workload") or "").lower()
    hx = _is_hx(fam, mtm)
    return dict(lines=lines, mtm=mtm, family=fam, plat=plat, nodes=int(n) if n else None, workload=wl,
                tce=bool(tce), raid=_num(raid if raid is not None else meta.get("raid")), allow_2node=allow_2node,
                is_hx=hx, is_nutanix=(wl == "nutanix") or hx, notes=list(meta.get("notes", [])),
                hx630v4=bool(re.match(r"7DG3", mtm) or re.search(r"HX630\s*V4", fam, I)),
                hx650v4=bool(re.match(r"7DG4", mtm) or re.search(r"HX650\s*V4", fam, I)))

def sel(c, rx, sec=None):
    rx = re.compile(rx, I) if isinstance(rx, str) else rx
    return [l for l in c["lines"] if (not sec or l["section"] == sec) and rx.search(l["descr"])]

qsum = lambda ls: sum(l["qty"] for l in ls)
ev = lambda ls: [f"{_fq(l['qty'])} x {l['fc']} {l['descr']}" for l in ls]

def cpu_lines(c):
    return [l for l in sel(c, r"(Xeon|EPYC).*Processor|Processor$", "Processor") if not re.search(r"Heatsink|Board|Clip|Module", l["descr"], I)]
def hs_lines(c):
    return [l for l in sel(c, r"Heatsink", "Processor") if not re.search(r"M\.2|SSD|Clip|Fan|Shroud|Bracket|Cable", l["descr"], I)]
def dimm_lines(c):
    return sel(c, r"\b\w*DIMM\b", "Memory")
def nic_lines(c):
    return sel(c, r"Ethernet Adapter")
def drive_lines(c):
    return [l for l in sel(c, r"\b(SSD|HDD)\b", "Storage") if not re.search(r"M\.2|Enablement|Adapter|Backplane|Filler|Tray|Interposer|HeatSink|Label|Cable|Kit\b", l["descr"], I)]
def nvme_drives(c):
    return [l for l in drive_lines(c) if re.search(r"NVMe", l["descr"], I)]
def dimm_gb(d):
    m = re.search(r"(\d+)\s*GB", d, I)
    return int(m.group(1)) if m else 0

def nic_vendor(d):
    if re.search(r"Mellanox|ConnectX|NVIDIA", d, I): return "Mellanox/NVIDIA"
    if re.search(r"Broadcom|57\d{3}|5719|5720", d, I): return "Broadcom"
    if re.search(r"Intel|E810|X710|I350|E610", d, I): return "Intel"
    if re.search(r"Marvell|QLogic", d, I): return "Marvell"
    return "other"

def nic_speed(d):
    m = re.search(r"(\d+)\s*/?\s*(\d+)?GbE|(\d+)GBASE-T", d, I)
    return "unknown" if not m else (m.group(2) or m.group(1) or m.group(3)) + "GbE"

def cap_tb(d):
    m = re.search(r"(\d+(?:\.\d+)?)TB", d, I)
    if m: return float(m.group(1))
    m = re.search(r"(\d+)GB", d, I)
    return int(m.group(1)) / 1000 if m else 0.0

def to_tb(value, unit):
    """Decimal TB, never TiB. Convert ONLY values tagged TiB (Nutanix Sizer, DCSC 'Node Tebibytes');
    drive marketing capacities are already decimal TB. Untagged units raise: tag at ingest."""
    v = float(value)
    u = str(unit or "").upper()
    tib = 2 ** 40 / 10 ** 12
    if u == "TB": return v
    if u == "TIB": return v * tib
    if u == "GB": return v / 1000
    if u == "GIB": return v * tib / 1024
    raise ValueError(f'to_tb: untagged unit "{unit}" - tag at ingest, never guess')

# ---------------------------------------------------------------- rule registry
RULES = []

def rule(rid, sev, applies=None):
    def deco(fn):
        RULES.append(dict(id=rid, severity=sev, applies=applies or (lambda c: True), fn=fn,
                          doc=(fn.__doc__ or "").strip().splitlines()[0] if fn.__doc__ else ""))
        return fn
    return deco

# ---- NICs -----------------------------------------------------------------------------
@rule("hx-nic-single-vendor", "block", lambda c: c["is_hx"])
def _(c):
    """HX nodes take ONE NIC vendor (one driver/firmware stack under Nutanix LCM)."""
    nl = nic_lines(c)
    v = sorted({nic_vendor(l["descr"]) for l in nl})
    if len(v) > 1:
        return [(f"ThinkAgile node mixes NIC vendors ({' + '.join(v)}) - single vendor required on HX (Nutanix LCM: one driver/firmware stack). Pair OCP and PCIe adapters from the same family.", nl)]

@rule("thinksystem-nic-mix-role", "warn", lambda c: not c["is_hx"])
def _(c):
    """ThinkSystem: two NIC vendors at the same speed/role is worth a second look (across roles is fine)."""
    by = collections.defaultdict(set)
    for l in nic_lines(c):
        by[nic_speed(l["descr"])].add(nic_vendor(l["descr"]))
    return [(f"Two NIC vendors at the same speed/role ({s}: {' + '.join(sorted(v))}) - fine across roles (data vs management), check it is deliberate here.", nic_lines(c))
            for s, v in by.items() if len(v) > 1]

# ---- HX platform rules ----------------------------------------------------------------
@rule("hx630-dual-socket", "block", lambda c: c["hx630v4"])
def _(c):
    """HX630 V4 needs exactly 2 populated sockets."""
    q = qsum(cpu_lines(c))
    if q != 2:
        return [(f"HX630 V4 requires exactly 2 populated sockets - config has {_fq(q)} CPU(s). Use 2x half-core-count CPUs to keep total cores (and Nutanix licensing) flat.", cpu_lines(c))]

@rule("hx630-384gb-unbuildable", "block", lambda c: c["hx630v4"])
def _(c):
    """HX630 V4 cannot build 384GB total memory (allowed DIMM counts skip it)."""
    dl = dimm_lines(c)
    if sum(dimm_gb(l["descr"]) * l["qty"] for l in dl) == 384:
        return [("HX630 V4 cannot build 384GB total memory (allowed DIMM counts skip it) - quote 256GB or 512GB.", dl)]

@rule("hx-nvme-single-sku", "block", lambda c: c["is_hx"] and not re.search(r"CTO2WW$", c["mtm"]))
def _(c):
    """HX all-flash NVMe drives are a single-SKU section: one drive part number per node."""
    nv = nvme_drives(c)
    sk = sorted({l["fc"] for l in nv})
    if len(sk) > 1:
        return [(f"HX all-flash NVME_DRIVES takes ONE drive SKU - config has {len(sk)} ({', '.join(sk)}). Collapse to a single capacity covering the total; uniform RI is the standard all-flash Nutanix build.", nv)]

@rule("hx650-storage-hybrid-floor", "block", lambda c: bool(re.match(r"7DG4CTO2WW$", c["mtm"])) or bool(re.search(r"HX650\s*V4\s*Storage", c["family"], I)))
def _(c):
    """HX650 V4 Storage hybrid: NVMe cache forces the rear NVMe backplane, so HDD count must be 10 or 12."""
    hd = [l for l in drive_lines(c) if re.search(r"HDD", l["descr"], I)]
    n = qsum(hd)
    if n and n not in (10, 12):
        return [(f"HX650 V4 Storage hybrid HDD count must be 10 or 12 per node (config has {_fq(n)}) - the rear NVMe cache backplane forces the floor. A 4-HDD hybrid will not validate; quote 10x or move to all-flash.", hd)]

@rule("hx650-bays-socket-min", "warn", lambda c: bool(re.match(r"7DG4CTO1WW$", c["mtm"])))
def _(c):
    """HX650 V4 with more than 8 NVMe drives on 1 CPU (16/24-bay territory). UNSETTLED, verify."""
    n, cp = qsum(nvme_drives(c)), qsum(cpu_lines(c))
    if n > 8 and cp < 2:
        return [(f"{_fq(n)} NVMe drives with {_fq(cp)} CPU - 16/24-bay HX650 V4 configs are believed to need 2 sockets (UNSETTLED: the evidence is correlational and the guide says 1 or 2 processors are supported; verify in the DCSC panel).", nvme_drives(c) + cpu_lines(c))]

@rule("nutanix-min-3-policy", "warn", lambda c: c["is_nutanix"] and c["nodes"] and c["nodes"] < 3 and not c["allow_2node"])
def _(c):
    """Policy: quote Nutanix clusters at 3 nodes minimum (--allow-2node to accept 2-node + witness)."""
    return [f"Policy: Nutanix clusters are quoted at 3 nodes minimum by default (config has {c['nodes']}). 2-node with a Witness VM is technically supported; pass --allow-2node to accept it."]

@rule("nutanix-2node-platform-truth", "info", lambda c: c["is_nutanix"] and c["nodes"] == 2)
def _(c):
    """Platform note: 2-node Nutanix exists with a Witness VM in a separate failure domain."""
    return ["Platform note: a 2-node Nutanix cluster is supported WITH a Witness VM in a separate failure domain (same resiliency as 3-node RF2; node removal unsupported)."]

@rule("nutanix-2node-capacity-cap", "warn", lambda c: c["is_nutanix"] and c["nodes"] == 2)
def _(c):
    """HX630/650 V4 guides: a 2-node configuration is limited to 20TB of storage per node."""
    raw = sum(cap_tb(l["descr"]) * l["qty"] for l in drive_lines(c))
    if raw > 20:
        return [(f"Per-node raw capacity is ~{raw:.0f}TB on a 2-node cluster - the HX630 V4 / HX650 V4 guides cap a 2-node configuration at 20TB of storage per node.", drive_lines(c))]

@rule("nutanix-capacity-gate", "warn", lambda c: c["is_nutanix"])
def _(c):
    """Per-node capacity above 307TB is only supported for Nutanix Unified Storage (hybrid band 307-550TB)."""
    dl = drive_lines(c)
    raw = sum(cap_tb(l["descr"]) * l["qty"] for l in dl)
    if raw <= 307:
        return []
    band = "over 550TB" if raw > 550 else "over 307TB"
    return [(f"Per-node raw capacity is ~{raw:.0f}TB ({band}). Lenovo Press (HX650 V4 guide): per-node capacity over 307TB is only supported for Nutanix Unified Storage use cases (hybrid 307TB to 550TB). Configurator messages can cite other thresholds on other configs - confirm the exact gate in the panel for this MTM.", dl)]

# ---- TCE ------------------------------------------------------------------------------
_ABSENT = [  # families with no Top Choice parts seen in validated exports as of 2026-09 (the list rotates)
    ("SR630 V3", r"\bSR630\s*V3\b", r"^7D73"), ("SR650 V3", r"\bSR650\s*V3\b", r"^7D76"),
    ("ST650 V3", r"\bST650\s*V3\b", r"^7D7A"), ("HX665 V3", r"\bHX665\s*V3\b", r"^7D9N"),
    ("HX630 V3", r"\bHX630\s*V3\b", r"^7D6MCTO3"),
    # gen1 HX has no version suffix; the lookahead stops a bare HX630 from swallowing HX630 V4
    ("HX630 (gen1)", r"\bHX630\b(?!\s*V\d)", r"^7D6MCTO4"), ("HX650 (gen1)", r"\bHX650\b(?!\s*V\d)", r"^7D6NCTO4"),
]
_PROVEN = re.compile(r"\b(ST250\s*V3|SR250\s*V3|SR635\s*V3|SR645\s*V3|SR655\s*V3|SR665\s*V3|SR630\s*V4|SR650a\s*V4|SR650\s*V4|HX630\s*V4|HX650\s*V4|FX630\s*V4|DM3010H|DE4200H)\b", I)
_OLDGEN = re.compile(r"\bV[12]\b", I)

def _absent(c):
    for label, fr, mr in _ABSENT:
        if re.search(fr, c["family"], I) or (c["mtm"] and re.search(mr, c["mtm"], I)):
            return label
    return None

_lbl = lambda c: c["family"] or c["mtm"] or "this platform"

@rule("tce-platform-support", "block", lambda c: c["tce"])
def _(c):
    """TCE on a platform with no Top Choice parts seen (SR630/650 V3, ST650 V3, HX665 V3, old HX) or V1/V2."""
    hit = _absent(c)
    if hit:
        return [f"TCE requested on {_lbl(c)} - no Top Choice parts have been seen on this platform (BU1E never appeared in validated configs as of 2026-09; the list rotates). Move to the V4 successor or drop the TCE requirement."]
    if _OLDGEN.search(c["family"]):
        return [f"TCE requested on {_lbl(c)} - Top Choice Express does not cover V1/V2 generation platforms. Move to a current generation or drop the TCE requirement."]

@rule("tce-platform-unproven", "warn", lambda c: c["tce"])
def _(c):
    """TCE on a platform never seen carrying BU1E: not the same as unavailable, confirm in the panel."""
    if _PROVEN.search(c["family"]) or _absent(c) or _OLDGEN.search(c["family"]):
        return []
    return [f"No Top Choice evidence either way for {_lbl(c)} - it has not been seen carrying BU1E in validated exports, which is NOT the same as unavailable. Confirm BU1E appears in the DCSC panel before quoting a fast-ship date."]

@rule("tce-all-or-nothing", "block", lambda c: c["tce"])
def _(c):
    """TCE is all-or-nothing: one non-TCE hardware line, a tied non-TCE order or a rack voids fast ship."""
    out = []
    bad = [l for l in c["lines"] if not is_service(l) and (l["tce"] is False or tce_verdict(l["fc"], c["mtm"]) is False)]
    if bad:
        out.append((f"TCE requested but {len(bad)} hardware line(s) are known non-TCE ({', '.join(l['fc'] for l in bad)}) - TCE is all-or-nothing per config, and a tied order with ANY non-TCE product voids the whole order.", bad))
    racks = sel(c, r"Rack Cabinet|\b42U\b|\b48U\b")
    if racks:
        out.append(("A rack chassis in the configuration voids TCE - quote racks on a separate order.", racks))
    return out

@rule("tce-dimm-max-64", "block", lambda c: c["tce"])
def _(c):
    """64GB is the largest Top Choice RDIMM; 96/128GB DIMMs are not TCE."""
    big = [l for l in dimm_lines(c) if dimm_gb(l["descr"]) > 64]
    if big:
        return [(f"TCE requested with {', '.join(str(dimm_gb(l['descr'])) + 'GB' for l in big)} DIMMs - 64GB is the largest Top Choice RDIMM. Re-spec the capacity with 64GB or smaller sticks.", big)]

@rule("tce-needs-panel-check", "info", lambda c: c["tce"])
def _(c):
    """A BOM alone cannot prove TCE: only hand-verified verdicts are checked here."""
    return [f"TCE can only be partly checked from a BOM ({len(TCE_OVERRIDES)} hand-verified part verdicts, platform tiers, DIMM ceiling, racks). Confirm BU1E in the live DCSC panel; the TCE list rotates."]

# ---- storage --------------------------------------------------------------------------
@rule("raid545-capability", "block")
def _(c):
    """RAID 5/6 requested on a 545-8i: that controller is cacheless entry-level and has no RAID 5/6."""
    k = sel(c, r"RAID 545-8i")
    if k and c["raid"] in (5, 6):
        return [(f"RAID {_fq(c['raid'])} requested on a 545-8i - that controller has no RAID 5/6 (cacheless). Use the 940-8i 4GB Flash.", k)]

@rule("nvme-behind-sas-raid", "warn")
def _(c):
    """NVMe data drives with only a plain SAS/SATA RAID controller and no visible NVMe path."""
    nv = nvme_drives(c)
    sas = [l for l in sel(c, r"RAID .*(SAS/SATA|12Gb)") if not re.search(r"Tri-?Mode|NVMe", l["descr"], I)]
    path = sel(c, r"Tri-?Mode|VROC|NVMe.*(Backplane|Retimer|Switch)")
    if nv and sas and not path:
        return [("NVMe data drives with only a SAS/SATA RAID controller and no visible NVMe path (tri-mode/VROC/NVMe backplane) - NVMe never runs behind plain SAS/SATA RAID. Confirm the drive attachment in DCSC.", nv + sas)]

@rule("controller-ports-vs-bays", "block")
def _(c):
    """A controller's port count must cover the drives behind it (an 8i caps at 8 drives)."""
    sd = [l for l in drive_lines(c) if re.search(r"SAS|SATA", l["descr"], I) and not re.search(r"NVMe", l["descr"], I)]
    n = qsum(sd)
    ctl = sel(c, r"(RAID|HBA).*(\b\d+i\b)")
    ports = max([0] + [int(re.search(r"\b(\d+)i\b", l["descr"]).group(1)) for l in ctl])
    if n and ports and n > ports:
        return [(f"{_fq(n)} SAS/SATA drives behind a {ports}i controller - port count must cover the drives (an 8i caps at 8 on a 12-bay backplane; step up to the 16i or add the direct riser the 16i needs on 12x3.5\").", sd + ctl)]

@rule("boot-mirror", "block", lambda c: bool(_m2_drives(c)))
def _(c):
    """M.2 boot drives should be a mirrored pair (qty 2 per node)."""
    d = _m2_drives(c)
    q = qsum(d)
    if q == 2: return []
    if q == 1: return [("Single M.2 boot drive (qty 1) - the M.2 RAID kit mirrors a PAIR. Quote 2x or the node boots off one unprotected drive.", d)]
    if q > 4: return [(f"M.2 boot drive count is {_fq(q)} - that is a cluster total, not a per-node quantity. Divide by the node count before reading this config.", d)]
    return [(f"M.2 boot drives = {_fq(q)}/node - expected 2 for a mirror.", d)]

def _m2_drives(c):
    return [l for l in sel(c, r"M\.2.*(SSD|960GB|480GB|1\.92TB)") if not re.search(r"Adapter|Kit|Cable|Enablement|HeatSink|Tray|Interposer|Filler|Label|Module", l["descr"], I)]

@rule("m2-kit-without-drives", "block", lambda c: bool(sel(c, r"M\.2.*(RAID|Enablement Kit)")))
def _(c):
    """An M.2 RAID/enablement kit with no M.2 drives behind it is a dead line item."""
    if not _m2_drives(c):
        return [("An M.2 RAID/enablement kit is quoted with no M.2 drives behind it - add the boot pair or drop the kit.", sel(c, r"M\.2.*(RAID|Enablement Kit)"))]

@rule("backplane-presence", "warn", lambda c: bool(drive_lines(c)))
def _(c):
    """Data drives quoted with no backplane line."""
    if not sel(c, r"Backplane"):
        return [("Data drives are quoted with no backplane line - DCSC derives some backplanes, but verify the drive bays actually exist in this chassis.", drive_lines(c))]

def bp_bays(d):
    return sum(int(m.group(1)) for m in re.finditer(r"(\d+)\s*x\s*([23]\.5)", re.split(r"\bwith\b", d or "", flags=I)[0], I))

@rule("drives-vs-backplane-bays", "warn")
def _(c):
    """More front drives than the selected backplane(s) hold."""
    fd = [l for l in c["lines"] if l["section"] == "Storage" and re.search(r"\b[23]\.5\b", l["descr"])
          and re.search(r"ssd|hdd|nvme|sas|sata", l["descr"], I)
          and not re.search(r"backplane|m\.2|enablement|cable|config\b|filler|tray|label|interposer|adapter|kit\b|raid|hba", l["descr"], I)]
    bps = [l for l in c["lines"] if l["section"] == "Storage" and re.search(r"backplane", l["descr"], I)]
    bays = sum(bp_bays(l["descr"]) * max(1, l["qty"]) for l in bps)
    n = qsum(fd)
    if bays and n > bays:
        return [(f"{_fq(n)} front drives but the selected backplane{'s hold' if len(bps) > 1 else ' holds'} {_fq(bays)} bays - add/upsize the backplane in DCSC or reduce drives.", fd + bps)]

@rule("35in-drives-on-25in-platform", "warn", lambda c: bool(c["plat"]) and c["plat"]["drives"]["supports_3_5"] is False)
def _(c):
    """3.5-inch data drives on a platform whose guide lists no 3.5-inch front bays."""
    d = [l for l in drive_lines(c) if re.search(r"\b3\.5\b", l["descr"])]
    if d:
        return [(f"3.5-inch drive(s) on {c['plat']['name']}, which the guide ({c['plat']['source']}) lists with no 3.5-inch front bays - match capacity with 2.5-inch drives (check rear-bay options in the guide).", d)]

# ---- compute / memory / sockets (need platform-specs) -----------------------------------
has_plat = lambda c: bool(c["plat"])

@rule("cpu-sockets-max", "block", lambda c: has_plat(c) and bool(cpu_lines(c)))
def _(c):
    """More CPUs than the platform has sockets."""
    q, mx = qsum(cpu_lines(c)), c["plat"]["sockets"]["max"]
    if q > mx:
        return [(f"{_fq(q)} processors exceeds the {mx}-socket maximum of {c['plat']['name']} ({c['plat']['source']}).", cpu_lines(c))]

@rule("cpu-sockets-min", "block", lambda c: has_plat(c) and bool(cpu_lines(c)) and not c["hx630v4"])
def _(c):
    """Fewer CPUs than the platform's minimum populated sockets."""
    q, mn = qsum(cpu_lines(c)), c["plat"]["sockets"]["min"]
    if 0 < q < mn:
        return [(f"{_fq(q)} processor(s) but {c['plat']['name']} needs at least {mn} ({c['plat']['source']}).", cpu_lines(c))]

@rule("hci-min-cpus", "warn", lambda c: has_plat(c) and bool(cpu_lines(c)) and not c["hx630v4"])
def _(c):
    """Fewer CPUs than an HCI guide's stated node minimum."""
    q, mn = qsum(cpu_lines(c)), c["plat"].get("hci_min_cpus") or 0
    if mn > c["plat"]["sockets"]["min"] and 0 < q < mn:
        return [(f"{_fq(q)} processor(s) but the {c['plat']['name']} guide states a minimum of {mn} for the node ({c['plat']['source']}).", cpu_lines(c))]

@rule("heatsink-cpu-pairing", "warn", lambda c: bool(cpu_lines(c)) and bool(hs_lines(c)))
def _(c):
    """Heatsinks pair 1:1 with CPUs."""
    cq, hq = qsum(cpu_lines(c)), qsum(hs_lines(c))
    if cq != hq:
        return [(f"{_fq(hq)} heatsink(s) for {_fq(cq)} processor(s) - heatsinks pair 1:1 with CPUs.", hs_lines(c) + cpu_lines(c))]

@rule("dimm-slots-exceeded", "block", lambda c: has_plat(c) and bool(dimm_lines(c)))
def _(c):
    """More DIMMs than the platform's slots (per populated socket when the CPU count is known)."""
    m, cp = c["plat"]["memory"], qsum(cpu_lines(c))
    n = qsum(dimm_lines(c))
    lim = m["dimm_slots_total"]
    why = f"the {lim} memory slots of {c['plat']['name']}"
    if cp and cp < c["plat"]["sockets"]["max"] and m.get("dimm_slots_per_socket"):
        lim = m["dimm_slots_per_socket"] * cp
        why = f"the {lim} slots available with {_fq(cp)} processor(s) ({m['dimm_slots_per_socket']} per socket) on {c['plat']['name']}"
    if n > lim:
        return [(f"{_fq(n)} DIMMs exceeds {why} ({c['plat']['source']}).", dimm_lines(c))]

@rule("dimm-cpu-balance", "warn", lambda c: bool(dimm_lines(c)))
def _(c):
    """DIMM count should split evenly across the populated CPUs."""
    cp, n = qsum(cpu_lines(c)), qsum(dimm_lines(c))
    if cp > 1 and n and n % cp:
        return [(f"{_fq(n)} DIMMs does not split evenly across {_fq(cp)} processors - balance the memory per CPU in DCSC.", dimm_lines(c))]

@rule("winserver-16core-min", "block", lambda c: bool(sel(c, r"Windows Server.*(Standard|Datacenter|Std)")))
def _(c):
    """Windows Server Standard/Datacenter licenses at least 16 cores per server."""
    ls = sel(c, r"Windows Server.*(Standard|Datacenter|Std)")
    cores = sum((int(m.group(1)) if (m := re.search(r"(\d+)[- ]?Core", l["descr"], I)) else 0) * (l["qty"] or 1) for l in ls)
    if 0 < cores < 16:
        return [(f"Windows Server licensing covers {_fq(cores)} cores - the license minimum is 16 cores per server. Add packs to reach 16.", ls)]

# ---- power ----------------------------------------------------------------------------
_PSU = re.compile(r"Power Supply", I)
_CORD = re.compile(r"Line Cord|Jumper Cord|Rack Power Cable|Power Cord", I)

def _volts(d):
    return [int(re.sub(r"\D", "", m.group(0))) for m in re.finditer(r"(\d{2,3})\s*V(?![a-z])", d, I)]
_is_range = lambda d: bool(re.search(r"(\d{2,3})\s*-\s*(\d{2,3})\s*V", d, I))

def _psu_low(d):
    """True = takes 115V, False = 230V-only, None = no verdict (-48V DC, unreadable)."""
    if re.search(r"-?48V\s*DC|\bDC\b", d, I): return None
    v = _volts(d)
    return None if not v else any(x <= 127 for x in v)

def _cord_low(d):
    if _is_range(d): return False                        # rack-fed 100-250V: the PDU decides
    if re.search(r"NEMA\s*5-(15|20)P", d, I): return True       # 120V US plug
    if re.search(r"NEMA\s*6-(15|20)P|NEMA\s*L6-", d, I): return False
    v = _volts(d)
    return bool(v) and max(v) <= 127

_watts = lambda d: int((re.search(r"(\d{3,4})\s*W\b", d, I) or [0, 0])[1] or 0)

@rule("psu-cord-voltage", "block", lambda c: bool(sel(c, _PSU)) and bool(sel(c, _CORD)))
def _(c):
    """230V-only PSU paired with a 120V line cord: the node will not power on at a 120V outlet."""
    hi = [l for l in sel(c, _PSU) if _psu_low(l["descr"]) is False]
    lo = [l for l in sel(c, _CORD) if _cord_low(l["descr"])]
    if not hi or not lo: return []
    w = _watts(hi[0]["descr"])
    fix = {1100: " On AMD V3 the same-wattage mixed-mode part is BNFH (1100W 230V/115V Platinum v3).",
           750: " On AMD V3 the same-wattage mixed-mode part is BNFG (750W 230V/115V Platinum v3)."}.get(w, "")
    return [(f"{', '.join(l['fc'] for l in hi)} is a 230V-ONLY power supply ({hi[0]['descr']}) but the config ships {', '.join(l['fc'] for l in lo)} - a 120V line cord ({lo[0]['descr']}). The node will not power on at a 120V outlet. Either move to a 230V/115V mixed-mode PSU of the same wattage, or confirm the site is 230V and switch to a 250V cord." + fix, hi + lo)]

@rule("psu-high-line-site-check", "warn", lambda c: any(_psu_low(l["descr"]) is False for l in sel(c, _PSU)))
def _(c):
    """230V-only PSU on 100-250V rack jumpers: correct in a 230V hall, a dead node in a 120V one."""
    cords = sel(c, _CORD)
    if not cords or any(_cord_low(l["descr"]) for l in cords) or not any(_is_range(l["descr"]) for l in cords):
        return []
    psu = next(l for l in sel(c, _PSU) if _psu_low(l["descr"]) is False)
    return [(f"{psu['fc']} is 230V-only ({psu['descr']}) and the cords are rack jumpers rated 100-250V, so the PDU decides. Confirm the site feeds 230V; a 120V hall needs the 230V/115V mixed-mode PSU instead.", [psu] + cords)]

@rule("d4390-pdu-gate", "block", lambda c: bool(sel(c, r"D4390")) or bool(re.search(r"D4390", c["family"], I)))
def _(c):
    """D4390 JBOD with the v1 0U 60A delta PDU: a DCSC compatibility-matrix gate."""
    p = sel(c, r"0U.*60A.*Delta.*(v1|\bV1\b)|v1.*0U.*60A")
    if p:
        return [("D4390 paired with the v1 0U 60A delta PDU - DCSC's compatibility matrix gates this pairing. Use the current PDU revision.", p)]

# ---- 4-port OCP -------------------------------------------------------------------------
@rule("ocp-x16-upgrade", "warn", lambda c: bool(re.search(r"\bV4\b", c["family"] + " " + c["mtm"], I)) and bool(sel(c, r"OCP.*4[- ]?Port|4[- ]?Port.*OCP")))
def _(c):
    """V4: a 4-port OCP adapter needs the x16 OCP cable kit (base OCP slots are x8)."""
    if not sel(c, r"OCP Cable Kit|x16 OCP|OCP.*BANDWIDTH|Bandwidth.*Upgrade"):
        return [("A 4-port OCP adapter is quoted on a V4 platform with no x16 OCP cable kit. Both OCP slots are x8 by default and a 4-port 25GbE OCP wants x16, so DCSC will normally pull the kit in (C1YK on SR650 V4/SR630 V4). Either accept it, or drop to a 2-port OCP adapter (e.g. BN2T) which fits the base x8 slot with no kit.", sel(c, r"OCP.*4[- ]?Port|4[- ]?Port.*OCP"))]

# ---- hygiene ----------------------------------------------------------------------------
@rule("config-completeness", "block", lambda c: len(c["lines"]) > 3)
def _(c):
    """A config missing processor, memory or power supply is a draft, not a quote; lone PSU is flagged."""
    need = [("a processor", cpu_lines), ("memory", dimm_lines), ("a power supply", lambda x: sel(x, r"Power Supply|\bPSU\b"))]
    miss = [n for n, f in need if not f(c)]
    out = []
    if miss:
        out.append("Config has no " + ", no ".join(miss) + " - DCSC will not close a configuration missing a required category.")
    psu = sel(c, r"Power Supply|\bPSU\b")
    if qsum(psu) == 1:
        out.append(("Only 1 power supply - production nodes are quoted 1+1 redundant. Confirm this is deliberate.", psu))
    return out

@rule("placeholder-junk", "block")
def _(c):
    """Sizer/competitor marker rows ('Do not quote this SKU') must not survive into a Lenovo BOM."""
    j = [l for l in c["lines"] if re.search(r"do not quote", l["descr"], I) or re.search(r"do not quote", l["fc"], I)]
    if j:
        return [("Non-quotable placeholder line(s) survived into the config: " + "; ".join(f"{l['fc'] or '?'} \"{l['descr'][:40]}\"" for l in j) + ". Drop them before pasting into DCSC.", j)]

@rule("duplicate-codes", "warn", lambda c: len(c["lines"]) > 1)
def _(c):
    """Repeated hardware feature code in one section (services/labels legitimately repeat)."""
    seen, dup = set(), []
    for l in c["lines"]:
        if not l["fc"] or l["section"] not in ("Processor", "Memory", "Storage", "PCI", "Power"): continue
        if re.search(r"Month|Warranty|Support|Premier|Service|Publication|Label", l["descr"], I): continue
        if l["fc"] in seen and l["fc"] not in dup: dup.append(l["fc"])
        seen.add(l["fc"])
    if dup:
        return [(f"Repeated hardware feature code(s) in one section: {', '.join(dup)} - usually a merge artifact; confirm the quantity is not double-counted.", [l for l in c["lines"] if l["fc"] in dup])]

@rule("platform-unresolved", "info", lambda c: not c["plat"])
def _(c):
    """MTM/family not in platform-specs.json: spec-based rules (sockets, DIMM slots, 3.5-inch) skipped."""
    return [f"Platform not recognised (mtm='{c['mtm']}', family='{c['family']}'): sockets/DIMM-slot/3.5-inch checks skipped. Pass --mtm, or see 'python specs.py spec --list'."]

@rule("bom-load-note", "info", lambda c: bool(c["notes"]))
def _(c):
    """Notes from loading the BOM (cluster-total division, unparsed lines, mixed MTMs)."""
    return list(c["notes"])

# ---------------------------------------------------------------- runner
_ORDER = {"block": 0, "warn": 1, "info": 2}

def validate(bom, **kw):
    """Run every rule. bom = list of {fc, descr, qty}; kw = mtm, family, nodes, workload, tce, raid, allow_2node."""
    c = make_ctx(bom, **kw)
    out = []
    for r in RULES:
        try:
            if not r["applies"](c): continue
            for x in r["fn"](c) or []:
                msg, evid = (x, []) if isinstance(x, str) else x
                out.append(dict(rule=r["id"], severity=r["severity"], message=msg, evidence=ev(evid)))
        except Exception as e:
            out.append(dict(rule=r["id"], severity="warn", message=f"rule crashed: {type(e).__name__}: {e} - treat as unchecked, not as passed", evidence=[]))
    out.sort(key=lambda f: _ORDER[f["severity"]])
    return out

# ---------------------------------------------------------------- deterministic corrections
def apply_fixups(bom):
    """Corrections with exactly one right answer. Returns (lines, notes); rules flag, this mutates copies."""
    out, notes = [dict(l) for l in (norm_line(x) if "descr" not in x else x for x in bom)], []
    jr = re.compile(r"do not quote|not included|no\s+\w+\s+(selected|included)|-none\b|none_lenovo", I)
    junk = [l for l in out if jr.search(l["descr"]) or jr.search(l["fc"])]
    if junk:
        out = [l for l in out if l not in junk]
        notes.append(f"Dropped {len(junk)} non-quotable placeholder line(s): " + "; ".join(l["descr"][:34] for l in junk) + ".")
    merged, by = [], {}
    for l in out:
        if not l["fc"]:
            merged.append(l); continue
        if l["fc"] in by:
            by[l["fc"]]["qty"] += l["qty"]
            notes.append(f"Merged duplicate line(s) for {l['fc']} into a single quantity.")
        else:
            by[l["fc"]] = l; merged.append(l)
    out = merged
    if any(re.search(r"M\.2.*(RAID|Enablement Kit)", l["descr"], I) for l in out):
        d = [l for l in out if re.search(r"M\.2", l["descr"], I) and re.search(r"SSD", l["descr"], I)
             and not re.search(r"Adapter|Kit|Cable|Enablement|HeatSink|Tray|Interposer|Filler", l["descr"], I)]
        if len(d) == 1 and d[0]["qty"] == 1:
            d[0]["qty"] = 2
            notes.append(f"Boot drives set to 2x {d[0]['fc']} - the M.2 RAID kit mirrors a pair, and a single boot drive leaves the node unprotected.")
    return out, notes

def diff_boms(intended, resolved, nodes=1):
    """Intended BOM (per node) vs what DCSC resolved (export lines, cluster qty): CLEAN or DRIFT."""
    want = {l["fc"]: l for l in intended if l["fc"]}
    got = {l["fc"]: dict(l, qty=l["qty"] / (nodes or 1)) for l in resolved if l["fc"]}
    missing = [w for k, w in want.items() if k not in got]
    drift = [(w, got[k]) for k, w in want.items() if k in got and w["qty"] and got[k]["qty"] and abs(got[k]["qty"] - w["qty"]) > 0.001]
    derived_rx = re.compile(r"Filler|Cable|Label|Tray|Interposer|Clip|Air duct|Riser|Cage|Board|Bezel|sponge|Placement|Months|Top Choice|Fan Module|Power Cord|Line Cord|Rail|Heatsink Clip|Lock|Latch|indicator|EIA|Neptune|Duct", I)
    extra = [g for k, g in got.items() if k not in want]
    unexpected = [g for g in extra if not derived_rx.search(g["descr"])]
    return dict(verdict="CLEAN" if not (missing or drift or unexpected) else "DRIFT", missing=missing, qty_drift=drift,
                unexpected=unexpected, derived_count=len(extra) - len(unexpected), tce_flag_present="BU1E" in got)

# ---------------------------------------------------------------- CLI
def _opt(a, name, cast=str, flag=False):
    if name in a:
        i = a.index(name)
        if flag:
            del a[i]; return True
        v = a[i + 1]; del a[i:i + 2]; return cast(v)
    return False if flag else None

def _show(findings, ctx):
    n = collections.Counter(f["severity"] for f in findings)
    for sev in ("block", "warn", "info"):
        fs = [f for f in findings if f["severity"] == sev]
        if not fs: continue
        print(f"\n== {sev.upper()} ({len(fs)}) " + "=" * 50)
        for f in fs:
            print(f"[{f['rule']}] {f['message']}")
            for e in f["evidence"][:6]: print(f"      {e}")
            if len(f["evidence"]) > 6: print(f"      ... +{len(f['evidence']) - 6} more")
    print(f"\n{n['block']} block, {n['warn']} warn, {n['info']} info  |  mtm={ctx['mtm'] or '?'} family={ctx['family'] or '?'} "
          f"nodes={ctx['nodes'] or '?'} workload={ctx['workload'] or '-'} tce={'yes' if ctx['tce'] else 'no'} lines={len(ctx['lines'])}")

def main(argv):
    for s in (sys.stdout, sys.stderr):
        try: s.reconfigure(encoding="utf-8", errors="replace")
        except Exception: pass
    a = list(argv)
    if not a or a[0] in ("-h", "--help", "help"):
        print(__doc__); return 0
    cmd, a = a[0], a[1:]
    if cmd == "list":
        for r in RULES: print(f"{r['id']:<32}{r['severity']:<6} {r['doc']}")
        print(f"\n{len(RULES)} rules; TCE_OVERRIDES holds {len(TCE_OVERRIDES)} hand-verified part verdicts")
        return 0
    try:
        if cmd == "check":
            kw = dict(mtm=_opt(a, "--mtm"), nodes=_opt(a, "--nodes", int), workload=_opt(a, "--workload"),
                      raid=_opt(a, "--raid", int), tce=True if _opt(a, "--tce", flag=True) else None,
                      allow_2node=_opt(a, "--allow-2node", flag=True))
            cluster, js = _opt(a, "--cluster-totals", flag=True), _opt(a, "--json", flag=True)
            if len(a) != 1: print(__doc__); return 2
            bom = load_bom(a[0])
            if cluster and not bom.meta.get("cluster_totals"):
                bom.meta["cluster_totals"] = True
                n = kw["nodes"] or bom.meta.get("nodes")
                if not n: sys.exit("--cluster-totals needs --nodes N (or a qty on the MTM line)")
                bom = _finish(list(bom), dict(bom.meta, nodes=n))
            if kw["workload"] and kw["workload"].lower() not in WORKLOADS:
                print(f"note: workload '{kw['workload']}' not recognised (known: {', '.join(WORKLOADS)})", file=sys.stderr)
            ctx = make_ctx(bom, **kw)
            fs = validate(bom, **kw)
            if js: print(json.dumps(fs, indent=1))
            else: _show(fs, ctx)
            return 1 if any(f["severity"] == "block" for f in fs) else 0
        if cmd == "diff":
            n = _opt(a, "--nodes", int) or 1
            if len(a) != 2: print(__doc__); return 2
            d = diff_boms(load_bom(a[0]), load_bom(a[1]), n)
            print(f"{d['verdict']}  (derived lines tolerated: {d['derived_count']}, BU1E present: {d['tce_flag_present']})")
            for w in d["missing"]: print(f"  missing     {w['fc']:<8}{_fq(w['qty']):>4} x {w['descr'][:60]}")
            for w, g in d["qty_drift"]: print(f"  qty drift   {w['fc']:<8}{_fq(w['qty'])} -> {_fq(g['qty'])}  {w['descr'][:50]}")
            for g in d["unexpected"]: print(f"  unexpected  {g['fc']:<8}{_fq(g['qty']):>4} x {g['descr'][:60]}")
            return 0 if d["verdict"] == "CLEAN" else 1
    except (OSError, ValueError, IndexError) as e:
        print(f"error: {e}", file=sys.stderr); return 2
    print(__doc__); return 2

if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
